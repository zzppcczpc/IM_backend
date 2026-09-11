import csv
import os
from pathlib import Path


def _read_text(path: str) -> str:
    raw = Path(path).read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _parse_text(path: str) -> tuple[str, dict]:
    text = _read_text(path)
    return text, {"parser": "plain_text", "character_count": len(text)}


def _parse_csv(path: str) -> tuple[str, dict]:
    rows = []
    with open(path, "r", encoding="utf-8-sig", errors="replace", newline="") as source:
        for row in csv.reader(source):
            rows.append("\t".join(cell.strip() for cell in row))
    column_count = max((len(row.split("\t")) for row in rows), default=0)
    return "\n".join(rows), {
        "parser": "csv",
        "row_count": len(rows),
        "column_count": column_count,
    }


def _parse_xlsx(path: str) -> tuple[str, dict]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sections = []
    row_count = 0
    for worksheet in workbook.worksheets:
        sections.append(f"[工作表: {worksheet.title}]")
        for row in worksheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value) for value in row]
            if any(value.strip() for value in values):
                sections.append("\t".join(values))
                row_count += 1
        sections.append("")
    return "\n".join(sections), {
        "parser": "openpyxl",
        "sheet_count": len(workbook.worksheets),
        "row_count": row_count,
    }


def _parse_docx(path: str) -> tuple[str, dict]:
    from docx import Document

    document = Document(path)
    sections = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table_index, table in enumerate(document.tables, start=1):
        sections.append(f"[表格 {table_index}]")
        for row in table.rows:
            sections.append("\t".join(cell.text.strip() for cell in row.cells))
    return "\n".join(sections), {
        "parser": "python-docx",
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
    }


def _parse_pptx(path: str) -> tuple[str, dict]:
    from pptx import Presentation

    presentation = Presentation(path)
    sections = []
    shape_count = 0
    for slide_index, slide in enumerate(presentation.slides, start=1):
        sections.append(f"[幻灯片 {slide_index}]")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                sections.append(shape.text.strip())
                shape_count += 1
    return "\n".join(sections), {
        "parser": "python-pptx",
        "slide_count": len(presentation.slides),
        "text_shape_count": shape_count,
    }


def _parse_pdf(path: str) -> tuple[str, dict]:
    from pypdf import PdfReader

    reader = PdfReader(path)
    pages = [(page.extract_text() or "").strip() for page in reader.pages]
    return "\n\n".join(page for page in pages if page), {
        "parser": "pypdf",
        "page_count": len(reader.pages),
    }


def parse_knowledge_base_file(path: str, extension: str) -> tuple[str, dict]:
    """根据扩展名提取文本；返回文本和可供后续切分使用的元数据。"""
    extension = extension.lower()
    parsers = {
        ".txt": _parse_text,
        ".md": _parse_text,
        ".csv": _parse_csv,
        ".xlsx": _parse_xlsx,
        ".docx": _parse_docx,
        ".pptx": _parse_pptx,
        ".pdf": _parse_pdf,
    }
    parser = parsers.get(extension)
    if not parser:
        raise ValueError(f"暂不支持解析文件类型: {extension}")
    if not os.path.isfile(path):
        raise FileNotFoundError("原文件不存在")
    text, metadata = parser(path)
    if not text.strip():
        raise ValueError("文件中没有提取到可用文本")
    metadata["text_length"] = len(text)
    return text, metadata
